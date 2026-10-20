"""
Client for S&B Filters' two shared dealer workbooks.

S&B publish no per-dealer feed and have no API. What they do maintain is two Google Sheets shared
with their whole dealer network, and this client reads both through Google's own xlsx export
endpoint (one request per workbook returns every tab, which is why they are parsed as Excel rather
than per-tab CSV -- the CSV endpoint needs a gid per tab and those change).

  * the **pricing master** is the authoritative catalog: ``All Product Info`` (every SKU with
    pricing, dimensions, compliance text and install links), ``Fitment``, ``miva export`` (a legacy
    storefront dump carrying fitment for much older SKUs) and ``Notes`` (dealer terms, prose).
  * the **Shopify export** adds the HTML descriptions and product imagery, matched in by SKU.

Because both files are the same for every dealer, their URLs come from ``src.constants`` /
``settings`` rather than from a connection's credentials -- contrast Rough Country, where every
dealer has their own jobber URL. Overrides exist for the management command and for tests.

Stock is deliberately not here: S&B send it by email on request, and it reaches us as a per-company
file on our relay. See ``relay.py``.
"""
import logging
import os
import tempfile
import time
import typing
import urllib.error
import urllib.request

import openpyxl
from django.conf import settings

from src import constants as src_constants
from src.integrations.clients.sb import exceptions
from src.integrations.clients.sb import feed_spec

logger = logging.getLogger(__name__)

_LOG_PREFIX = "[SB-SHEET-CLIENT]"

DEFAULT_FILE_MAX_AGE_SECONDS = 6 * 60 * 60  # 6 hours
_DOWNLOAD_TIMEOUT_SECONDS = 300
_HEAD_TIMEOUT_SECONDS = 30

PRICING_MASTER_LOCAL_FILENAME = "sb_pricing_master.xlsx"
SHOPIFY_EXPORT_LOCAL_FILENAME = "sb_shopify_export.xlsx"

_REQUEST_HEADERS = {
    "User-Agent": "Mozilla/5.0 (compatible; AfterMarketScout/1.0; +https://aftermarketscout.com)",
    "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*",
}

# Google serves a login page (HTML, 200) rather than a 403 when a sheet stops being link-shared,
# so a content check is the only way to notice. xlsx is a zip: it starts with "PK".
_XLSX_MAGIC = b"PK"


def _sheet_export_url(sheet_id: str) -> str:
    return src_constants.SB_SHEET_EXPORT_URL_TEMPLATE.format(sheet_id=sheet_id)


def _setting_str(name: str, fallback: str) -> str:
    raw = getattr(settings, name, None)
    value = str(raw).strip() if raw is not None else ""
    return value or fallback


class SbSheetClient:
    """
    Downloads and parses S&B's pricing master and Shopify export.

    Both workbooks are cached on local disk and re-downloaded only when missing or older than
    ``file_max_age`` -- the same contract as the Rough Country client, and the reason the nightly
    ingest and a per-company pricing job in the same window don't each pull 2 MB of Google Sheets.
    """

    def __init__(
        self,
        pricing_master_url: typing.Optional[str] = None,
        shopify_export_url: typing.Optional[str] = None,
        local_dir: typing.Optional[str] = None,
        pricing_master_path: typing.Optional[str] = None,
        shopify_export_path: typing.Optional[str] = None,
        file_max_age: int = DEFAULT_FILE_MAX_AGE_SECONDS,
    ):
        self.pricing_master_url = (pricing_master_url or "").strip() or _setting_str(
            "SB_PRICING_MASTER_URL",
            _sheet_export_url(src_constants.SB_PRICING_MASTER_SHEET_ID),
        )
        self.shopify_export_url = (shopify_export_url or "").strip() or _setting_str(
            "SB_SHOPIFY_EXPORT_URL",
            _sheet_export_url(src_constants.SB_SHOPIFY_EXPORT_SHEET_ID),
        )
        self.local_dir = (local_dir or "").strip() or _setting_str(
            "SB_LOCAL_FEED_DIR", tempfile.gettempdir()
        )
        # Explicit local paths bypass downloading entirely -- management command / test use.
        self.pricing_master_path = (pricing_master_path or "").strip() or None
        self.shopify_export_path = (shopify_export_path or "").strip() or None
        self.file_max_age = file_max_age

    # ----------------------------------------------------------------------------------
    # Download
    # ----------------------------------------------------------------------------------
    def _local_path(self, filename: str) -> str:
        return os.path.join(self.local_dir, filename)

    def pricing_master_local_path(self) -> str:
        return self.pricing_master_path or self._local_path(PRICING_MASTER_LOCAL_FILENAME)

    def shopify_export_local_path(self) -> str:
        return self.shopify_export_path or self._local_path(SHOPIFY_EXPORT_LOCAL_FILENAME)

    def is_file_outdated(self, path: str) -> bool:
        if not os.path.exists(path):
            return True
        return (time.time() - os.path.getmtime(path)) > self.file_max_age

    def _download(self, url: str, path: str) -> str:
        """
        Fetch one workbook to ``path`` via a temp file, so a failed or truncated download never
        replaces a good cached copy with a broken one.
        """
        if not url:
            raise ValueError("S&B workbook URL is required.")

        tmp_path = "{}.tmp".format(path)
        try:
            logger.info("{} Downloading {} -> {}.".format(_LOG_PREFIX, url, path))
            request = urllib.request.Request(url, headers=_REQUEST_HEADERS)
            with urllib.request.urlopen(request, timeout=_DOWNLOAD_TIMEOUT_SECONDS) as response:
                payload = response.read()
        except urllib.error.HTTPError as e:
            raise exceptions.SbDownloadError(
                "HTTP {} downloading S&B workbook {}: {}.".format(e.code, url, e.reason)
            )
        except urllib.error.URLError as e:
            raise exceptions.SbDownloadError(
                "Could not reach S&B workbook {}: {}.".format(url, e.reason)
            )
        except Exception as e:
            raise exceptions.SbDownloadError(
                "Failed to download S&B workbook {}: {}.".format(url, str(e))
            )

        if not payload.startswith(_XLSX_MAGIC):
            # See _XLSX_MAGIC: a de-shared sheet answers 200 with an HTML sign-in page.
            raise exceptions.SbDownloadError(
                "S&B workbook {} did not return an xlsx file ({} bytes, starts {!r}). The sheet "
                "is most likely no longer shared with anyone who has the link.".format(
                    url, len(payload), payload[:16],
                )
            )

        try:
            os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
            with open(tmp_path, "wb") as handle:
                handle.write(payload)
            os.replace(tmp_path, path)
        except OSError as e:
            if os.path.exists(tmp_path):
                try:
                    os.remove(tmp_path)
                except OSError:
                    pass
            raise exceptions.SbDownloadError(
                "Failed to write S&B workbook to {}: {}.".format(path, str(e))
            )

        logger.info("{} Saved {} ({} bytes).".format(_LOG_PREFIX, path, len(payload)))
        return path

    def download(self, force_download: bool = False) -> typing.Dict[str, str]:
        """Ensure both workbooks are on disk and fresh; returns their paths."""
        paths = {}
        for key, url, path, explicit in (
            ("pricing_master", self.pricing_master_url, self.pricing_master_local_path(), self.pricing_master_path),
            ("shopify_export", self.shopify_export_url, self.shopify_export_local_path(), self.shopify_export_path),
        ):
            if explicit:
                if not os.path.exists(path):
                    raise exceptions.SbDownloadError(
                        "S&B {} local file not found: {}.".format(key, path)
                    )
                logger.info("{} Using local {} {}.".format(_LOG_PREFIX, key, path))
            elif force_download or self.is_file_outdated(path):
                self._download(url, path)
            else:
                logger.info("{} Using cached {} {}.".format(_LOG_PREFIX, key, path))
            paths[key] = path
        return paths

    def test_connection(self, timeout: int = _HEAD_TIMEOUT_SECONDS) -> None:
        """
        Confirm both workbooks are still reachable and still shared, without pulling 2 MB: asks
        for the first KB via a Range header. Servers that ignore Range just send more than we
        read. Raises SbDownloadError on anything that isn't a readable xlsx.
        """
        for url in (self.pricing_master_url, self.shopify_export_url):
            if not url:
                raise ValueError("S&B workbook URL is required.")
            headers = dict(_REQUEST_HEADERS)
            headers["Range"] = "bytes=0-1023"
            request = urllib.request.Request(url, headers=headers)
            try:
                with urllib.request.urlopen(request, timeout=timeout) as response:
                    status = response.status if hasattr(response, "status") else response.getcode()
                    chunk = response.read(1024)
            except urllib.error.HTTPError as e:
                raise exceptions.SbDownloadError(
                    "HTTP {} checking S&B workbook {}: {}.".format(e.code, url, e.reason)
                )
            except urllib.error.URLError as e:
                raise exceptions.SbDownloadError(
                    "Could not reach S&B workbook {}: {}.".format(url, e.reason)
                )

            if status not in (200, 206) or not chunk:
                raise exceptions.SbDownloadError(
                    "Unexpected response (status={}) checking S&B workbook {}.".format(status, url)
                )
            if not chunk.startswith(_XLSX_MAGIC):
                raise exceptions.SbDownloadError(
                    "S&B workbook {} is no longer returning a spreadsheet — check that it is "
                    "still shared with anyone who has the link.".format(url)
                )

    # ----------------------------------------------------------------------------------
    # Parse
    # ----------------------------------------------------------------------------------
    @staticmethod
    def _open_workbook(path: str) -> openpyxl.Workbook:
        try:
            return openpyxl.load_workbook(path, read_only=True, data_only=True)
        except Exception as e:
            raise exceptions.SbParseError(
                "Failed to open S&B workbook {}: {}.".format(path, str(e))
            )

    @staticmethod
    def _rows(worksheet: typing.Any) -> typing.List[typing.Tuple]:
        """
        Every row with at least one populated cell. ``read_only`` worksheets are generators, and
        both workbooks carry trailing blank rows from Google's export, so this both materializes
        and prunes.
        """
        return [
            row
            for row in worksheet.iter_rows(values_only=True)
            if row and any(cell is not None for cell in row)
        ]

    @staticmethod
    def _dicts_from_rows(
        rows: typing.Sequence[typing.Tuple],
        spec: typing.Mapping[str, typing.Sequence[str]],
        header_index: int = 0,
        require: typing.Sequence[str] = (),
    ) -> typing.List[typing.Dict[str, typing.Any]]:
        """
        Turn positional rows into ``{field: value}`` dicts using ``spec``. Rows are kept even when
        most fields are blank; dropping them is the service's call, not the parser's.
        """
        if header_index >= len(rows):
            return []
        columns = feed_spec.map_columns(rows[header_index], spec)
        missing = [field for field in require if field not in columns]
        if missing:
            raise exceptions.SbParseError(
                "S&B sheet is missing required column(s) {}; header was {!r}.".format(
                    ", ".join(missing), rows[header_index][:20],
                )
            )
        out = []
        for row in rows[header_index + 1:]:
            record = {}
            for field, index in columns.items():
                record[field] = row[index] if index < len(row) else None
            out.append(record)
        return out

    @classmethod
    def _catalog_header_index(cls, rows: typing.Sequence[typing.Tuple]) -> int:
        """
        Row index of the catalog header. S&B keep two note rows above it today, but the count has
        changed before, so it is found by looking for the ``sku`` column instead of assumed.
        """
        limit = min(len(rows), feed_spec.CATALOG_MAX_HEADER_SCAN_ROWS)
        for index in range(limit):
            columns = feed_spec.map_columns(rows[index], {"sku": feed_spec.CATALOG_HEADER_MARKER_COLUMNS})
            if "sku" in columns:
                return index
        raise exceptions.SbParseError(
            "Could not find the 'sku' header column in S&B's catalog tab within the first {} "
            "rows.".format(limit)
        )

    def _parse_pricing_master(self, path: str) -> typing.Dict[str, typing.Any]:
        workbook = self._open_workbook(path)
        try:
            sheet_names = list(workbook.sheetnames)

            catalog_tab = feed_spec.find_tab(sheet_names, feed_spec.CATALOG_TAB_KEYS)
            if not catalog_tab:
                raise exceptions.SbParseError(
                    "S&B pricing master has no catalog tab (looked for {!r} in {!r}).".format(
                        feed_spec.CATALOG_TAB_KEYS, sheet_names,
                    )
                )
            catalog_rows = self._rows(workbook[catalog_tab])
            header_index = self._catalog_header_index(catalog_rows)
            catalog = self._dicts_from_rows(
                catalog_rows,
                feed_spec.CATALOG_COLUMNS,
                header_index=header_index,
                require=("sku",),
            )

            # The fitment tab's name also contains "fitment" in the Shopify workbook, but this is
            # the pricing master, so a single match is unambiguous here.
            fitment: typing.List[typing.Dict[str, typing.Any]] = []
            fitment_tab = feed_spec.find_tab(sheet_names, feed_spec.FITMENT_TAB_KEYS)
            if fitment_tab:
                fitment = self._dicts_from_rows(
                    self._rows(workbook[fitment_tab]), feed_spec.FITMENT_COLUMNS, require=("sku",)
                )
            else:
                logger.warning("{} No Fitment tab in {!r}.".format(_LOG_PREFIX, sheet_names))

            miva: typing.List[typing.Dict[str, typing.Any]] = []
            miva_tab = feed_spec.find_tab(sheet_names, feed_spec.MIVA_TAB_KEYS)
            if miva_tab:
                miva = self._dicts_from_rows(self._rows(workbook[miva_tab]), feed_spec.MIVA_COLUMNS)
            else:
                logger.info("{} No miva export tab in {!r}.".format(_LOG_PREFIX, sheet_names))

            notes_tab = feed_spec.find_tab(sheet_names, feed_spec.NOTES_TAB_KEYS)
            notes_rows = len(self._rows(workbook[notes_tab])) if notes_tab else 0
        finally:
            workbook.close()

        logger.info(
            "{} Pricing master: catalog={} fitment={} miva={} notes_rows={} (tabs={!r}).".format(
                _LOG_PREFIX, len(catalog), len(fitment), len(miva), notes_rows, sheet_names,
            )
        )
        return {"catalog": catalog, "fitment": fitment, "miva": miva, "sheet_names": sheet_names}

    def _parse_shopify_export(self, path: str) -> typing.List[typing.Dict[str, typing.Any]]:
        """
        Rows from every Shopify tab, in SHOPIFY_TAB_PRIORITY order with a ``tab`` field added.

        Order is the whole point: the two tabs are the same export at different dates, and the
        service keeps the first value it sees per SKU, so the current tab must come first. Rows
        stay in sheet order within a tab because Shopify's format is positional -- a product's
        extra image rows follow its first row and carry nothing else.
        """
        workbook = self._open_workbook(path)
        try:
            sheet_names = list(workbook.sheetnames)
            tabs = feed_spec.find_tabs(sheet_names, feed_spec.SHOPIFY_TAB_PRIORITY)
            if not tabs:
                raise exceptions.SbParseError(
                    "S&B Shopify export has no recognizable tab (looked for {!r} in {!r}).".format(
                        feed_spec.SHOPIFY_TAB_PRIORITY, sheet_names,
                    )
                )
            out: typing.List[typing.Dict[str, typing.Any]] = []
            for tab in tabs:
                rows = self._dicts_from_rows(
                    self._rows(workbook[tab]), feed_spec.SHOPIFY_COLUMNS, require=("product_id",)
                )
                for row in rows:
                    row["tab"] = tab
                out.extend(rows)
                logger.info("{} Shopify tab {!r}: {} rows.".format(_LOG_PREFIX, tab, len(rows)))
        finally:
            workbook.close()
        return out

    def get_feed_data(
        self, download_if_missing: bool = True, force_download: bool = False
    ) -> typing.Dict[str, typing.Any]:
        """
        Both workbooks parsed into row dicts: ``catalog``, ``fitment``, ``miva``, ``shopify``.

        Shaping those into parts and fitments is the service's job (see
        ``src.integrations.services.sb``) -- this stays a parser so the grouping rules are
        testable without a workbook.
        """
        if download_if_missing or force_download:
            paths = self.download(force_download=force_download)
        else:
            paths = {
                "pricing_master": self.pricing_master_local_path(),
                "shopify_export": self.shopify_export_local_path(),
            }
            for key, path in paths.items():
                if not os.path.exists(path):
                    raise exceptions.SbDownloadError(
                        "S&B {} file not found and downloading is disabled: {}.".format(key, path)
                    )

        master = self._parse_pricing_master(paths["pricing_master"])
        shopify = self._parse_shopify_export(paths["shopify_export"])

        logger.info(
            "{} Loaded catalog={} fitment={} miva={} shopify={}.".format(
                _LOG_PREFIX,
                len(master["catalog"]),
                len(master["fitment"]),
                len(master["miva"]),
                len(shopify),
            )
        )
        return {
            "catalog": master["catalog"],
            "fitment": master["fitment"],
            "miva": master["miva"],
            "shopify": shopify,
        }
